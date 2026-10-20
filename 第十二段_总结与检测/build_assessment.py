from pathlib import Path
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parent;OUT=R.parent/'demooo/独立片段/12_总结与检测/03_评价附件_结尾情境题.docx';OUT.parent.mkdir(parents=True,exist_ok=True)
d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)
for n in ['Normal','Title','Heading 1','Heading 2']:
 s=d.styles[n];s.font.name='Microsoft YaHei';s._element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑');s.font.color.rgb=RGBColor(0,0,0);s.font.size=Pt(10.5 if n=='Normal' else 18 if n=='Title' else 12);s.paragraph_format.space_after=Pt(5);s.paragraph_format.space_before=Pt(8 if n.startswith('Heading') else 0)
 for border in list(s._element.findall('.//'+qn('w:pBdr'))):border.getparent().remove(border)
d.styles['Normal'].paragraph_format.line_spacing=1.16
def p(s='',style=None):return d.add_paragraph(s,style)
def h(s):d.add_heading(s,level=1)
d.add_heading('海洋发烧了结尾情境评价',0)
p('03号评价附件　学生作答页')
p('本题用于检查海温距平判读、海气作用解释、区域迁移和应对决策。视频中有3秒思考停顿；正式作答建议暂停视频，用3—5分钟完成。满分4分。')
p('姓名 ____________________　班级 ______________　日期 ______________')
h('教学情境')
p('以下为原创教学假设，不是历史事件记录，也不是当年真实预报：Niño3.4区持续异常偏暖，信风减弱。石河子某校收到当地强降雨预警。')
p('说明：题目已给出异常偏暖的条件，不要求仅凭一条海温信息诊断厄尔尼诺事件；也未给出预警等级、时段或确定性，因此不得擅自补充。')
h('作答要求')
qs=['海温距平应为正值还是负值？请写出判断依据。','与正常状态相比，赤道东太平洋冷水上升流怎样变化？请说明信风与上升流之间的联系。','能否仅凭厄尔尼诺信号，断定石河子一定发生暴雨？为什么？','依据当地强降雨预警，为学校提出一项可执行行动，并判断它属于适应还是减缓。']
for i,q in enumerate(qs):
 p(f'{i+1}．{q}');p('答：____________________________________________________________');p('_______________________________________________________________')
teacher_title=d.add_heading('参考答案与评分标准',0);teacher_title.paragraph_format.page_break_before=True;p('03号评价附件　教师评分页')
p('四个维度各1分，总分4分。允许使用含义相同的表述；评价重点是证据、联系和行动类别，而不是背诵固定措辞。')
answers=[('1 海温判读','正值。海温距平＝当前海温－常年同期平均海温；题设异常偏暖，差值为正。','正负判断正确0.5分；说明与常年同期比较或给出正确计算关系0.5分。'),('2 因果解释','东部冷水上升流减弱。信风减弱，表层暖水向西堆积的作用减弱，暖水向东扩展；东部温跃层加深，上升的冷水对表层的冷却作用减弱。','指出上升流减弱0.5分；正确联系信风减弱与暖水分布或东部温跃层变化0.5分。'),('3 区域迁移','不能仅凭厄尔尼诺信号断定石河子一定暴雨。远距离气候联系具有区域、季节和事件差异，具体风险须结合本地气象监测、预警与其他天气条件。','明确不能作必然推断0.5分；指出需要本地证据或区域季节差异0.5分。'),('4 应对决策','例如：按学校安全预案和当地预警指引，及时调整或暂停户外活动，安排安全场所并提醒学生避开积水危险区域。这属于适应。','提出符合预警指引、可执行的安全行动0.5分；正确归为适应0.5分。')]
for name,answer,score in answers:
 h(name);p(answer);p('计分：'+score)
h('使用边界')
p('本地预警是应对决策的直接依据，不能将本题写成厄尔尼诺直接造成石河子暴雨的证据。预警等级与有效期应以实际官方发布为准，不自动宣布全校停课。关掉不用的灯属于减缓的例子，但不能替代强降雨预警下的及时安全行动。')
h('科学依据')
p('NOAA Ocean Service：What are El Niño and La Niña；NOAA JetStream：Effects of ENSO in the Pacific。中国气象局：气象灾害预警信号发布与传播管理办法。完整来源链接见本段制作依据与来源文件。')
footer=sec.footer.paragraphs[0];footer.alignment=2
run=footer.add_run('海洋发烧了　03号评价附件　');run.font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
d.save(OUT);print(OUT)
